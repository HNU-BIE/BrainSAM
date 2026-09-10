import json
import os
from pathlib import Path

import openpyxl

from inference.click_correction import (
    seg_nifti_inference_with_correction,
    MAX_CORRECTION_ROUNDS,
)
from config import get_args


def load_species_paths(json_path: str) -> dict:
    with open(json_path, "r") as f:
        data = json.load(f)

    def read_and_join(folder: str, txt_file: str) -> list[str]:
        with open(txt_file, "r", encoding="utf-8-sig") as f:
            filenames = f.read().strip().splitlines()
        return [str(Path(folder) / name.strip()) for name in filenames if name.strip()]

    return {
        species: (
            read_and_join(paths[0], paths[2]),
            read_and_join(paths[1], paths[3]),
        )
        for species, paths in data.items()
    }


def build_header(max_rounds: int) -> list:
    """Case, Dice_r0/ASD_r0/HD95_r0, Dice_r1/..., ..., rounds_used, corrected_mask_path"""
    header = ["Case"]
    for r in range(max_rounds + 1):
        header += [f"Dice_r{r}", f"ASD_r{r}", f"HD95_r{r}"]
    header += ["rounds_used", "corrected_mask_path"]
    return header


def init_xlsx(xlsx_path: str, sheet_name: str, header: list):
    if os.path.exists(xlsx_path):
        wb = openpyxl.load_workbook(xlsx_path)
        if sheet_name not in wb.sheetnames:
            ws = wb.create_sheet(title=sheet_name)
            ws.append(header)
            wb.save(xlsx_path)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        ws.append(header)
        wb.save(xlsx_path)


def append_to_xlsx(xlsx_path: str, sheet_name: str, rows: list):
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb[sheet_name]
    for row in rows:
        ws.append(row)
    wb.save(xlsx_path)


def load_progress_from_log(log_path: str) -> dict:
    if os.path.exists(log_path):
        with open(log_path, "r") as f:
            data = json.load(f)
        print(f"[Resume] 从日志恢复 {len(data)} 条已完成记录: {log_path}")
        return data
    return {}


def save_progress_to_log(log_path: str, results_dict: dict):
    with open(log_path, "w") as f:
        json.dump(results_dict, f, indent=2)


def build_row(case_name: str, entry: dict, max_rounds: int) -> list:
    """把一个 case 的纠正结果（或 FAILED 标记）展开成一行 xlsx 数据，
    round 数不足 max_rounds+1 时（比如提前 early-stop）用空字符串补齐对齐列。"""
    if entry.get("error"):
        row = [case_name]
        row += ["FAILED", entry["error"], ""]
        row += [""] * (3 * max_rounds)  # 剩余轮次列留空
        row += ["", ""]
        return row

    round_results = entry["round_results"]
    row = [case_name]
    for r in range(max_rounds + 1):
        if r < len(round_results):
            rr = round_results[r]
            row += [f"{rr['dice']:.4f}", f"{rr['asd']:.4f}", f"{rr['hd95']:.4f}"]
        else:
            row += ["", "", ""]
    row += [entry["rounds_used"], entry["corrected_save_path"]]
    return row


def main():
    from config import get_args, EXPERIMENTS
    from inference.base_inference import build_brainsam_predictor

    # 模型只加载一次（用第0个config初始化，ckpt/device所有实验相同）
    args0 = get_args(0)
    predictor = build_brainsam_predictor(
        config_file=args0.model_cfg,
        ckpt_path=args0.ckpt,
        device=args0.device,
    )

    header = build_header(MAX_CORRECTION_ROUNDS)

    for exp_idx in range(len(EXPERIMENTS)):
        args = get_args(exp_idx)

        print(f"\n{'='*60}")
        print(f"[EXP {exp_idx+1}/{len(EXPERIMENTS)}] species={args.species}  sheet={args.sheet}")
        print(f"  modal_id={args.modal_id}  organ_ids={args.organ_ids}")
        print(f"  save_dir={args.save_dir}")
        print(f"{'='*60}")

        xlsx_path  = args.xlsx
        sheet_name = args.sheet
        log_dir    = args.log_dir or args.save_dir
        log_path   = os.path.join(log_dir, f"checkpoint_{sheet_name}.json")

        os.makedirs(args.save_dir, exist_ok=True)
        os.makedirs(log_dir, exist_ok=True)

        file_dataset = load_species_paths(args.config_json)
        raw_path_list, gt_path_list = file_dataset[args.species]

        init_xlsx(xlsx_path, sheet_name, header)
        results_dict    = load_progress_from_log(log_path)
        completed_cases = set(results_dict.keys())

        pending = [
            (nii, gt)
            for nii, gt in zip(raw_path_list, gt_path_list)
            if os.path.basename(nii) not in completed_cases
        ]
        print(f"待处理: {len(pending)} / {len(raw_path_list)} cases  (已跳过: {len(completed_cases)})")

        new_since_last_save = []

        for i, (nii_path, gt_path) in enumerate(pending):
            case_name = os.path.basename(nii_path)
            subj     = os.path.basename(nii_path).replace(".nii.gz", "").replace(".nii", "")
            corrected_out_path = os.path.join(args.save_dir, f"{subj}_mask_corrected.nii.gz")
            if os.path.exists(corrected_out_path):
                print(f"\n  [{i+1}/{len(pending)}] {subj}  [SKIP] corrected mask already exists")
                continue
            try:
                result = seg_nifti_inference_with_correction(
                    predictor,
                    video_dir=nii_path,
                    ptsList=None,
                    ptsTypeList=None,
                    axis=args.axis,
                    save_dir=args.save_dir,
                    nii_idx="Autonomous",
                    auto_prompt=args.auto_prompt,
                    modal_id=args.modal_id,
                    organ_ids=args.organ_ids,
                    gt_path=gt_path,
                )
                dice0 = result["round_results"][0]["dice"]
                dice_final = result["round_results"][-1]["dice"]
                print(f"  Case: {case_name}  Dice_r0: {dice0:.4f} -> Dice_final: {dice_final:.4f}  "
                      f"(rounds_used={result['rounds_used']})")
                results_dict[case_name] = result
                new_since_last_save.append(build_row(case_name, result, MAX_CORRECTION_ROUNDS))
            except Exception as e:
                print(f"  [ERROR] {case_name}: {e}")
                results_dict[case_name] = {"error": str(e)}
                new_since_last_save.append(build_row(case_name, {"error": str(e)}, MAX_CORRECTION_ROUNDS))
            print("  " + "-"*50)

            if (i + 1) % args.save_interval == 0:
                append_to_xlsx(xlsx_path, sheet_name, new_since_last_save)
                save_progress_to_log(log_path, results_dict)
                new_since_last_save = []

        if new_since_last_save:
            append_to_xlsx(xlsx_path, sheet_name, new_since_last_save)
            save_progress_to_log(log_path, results_dict)

        print(f"[EXP {exp_idx+1}] Done. Results -> {xlsx_path} [{sheet_name}]")

    print("\n[ALL DONE]")


if __name__ == "__main__":
    main()